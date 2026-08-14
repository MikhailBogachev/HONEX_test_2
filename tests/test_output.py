"""Tests for output writers — CSV, JSON, XLSX consistency."""

from __future__ import annotations

import csv
import json

import pytest

from src.models import Listing, Seller, SellerType, EnrichmentStatus, DataOrigin, RunSummary


@pytest.fixture
def sample_listings(make_listing, make_seller):
    return [
        make_listing(
            avito_id="100",
            title="Фазорегулятор 243502B010 Новое",
            price=4500.0,
            seller=make_seller(seller_name="Продавец A"),
        ),
        make_listing(
            avito_id="101",
            title="Фазорегулятор 243502B010 оригинал Новое",
            price=5200.0,
            seller=make_seller(seller_id="seller_002", seller_name="Продавец B"),
        ),
    ]


class TestCSV:
    def test_csv_creates_file(self, tmp_path, sample_listings):
        from src.output.output import write_csv

        path = tmp_path / "test.csv"
        write_csv(sample_listings, path)

        assert path.exists()

    def test_csv_has_header_and_rows(self, tmp_path, sample_listings):
        from src.output.output import write_csv

        path = tmp_path / "test.csv"
        write_csv(sample_listings, path)

        with open(path, encoding="utf-8") as f:
            reader = csv.DictReader(f)
            rows = list(reader)

        assert len(rows) == 2
        assert "source_article" in rows[0]
        assert "seller_name" in rows[0]
        assert rows[0]["title"] == "Фазорегулятор 243502B010 Новое"

    def test_csv_empty_list(self, tmp_path):
        from src.output.output import write_csv

        path = tmp_path / "empty.csv"
        write_csv([], path)

        assert path.exists()
        with open(path, encoding="utf-8") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
        assert len(rows) == 0


class TestJSON:
    def test_json_creates_file(self, tmp_path, sample_listings):
        from src.output.output import write_json

        path = tmp_path / "test.json"
        write_json(sample_listings, path)

        assert path.exists()

    def test_json_has_correct_structure(self, tmp_path, sample_listings):
        from src.output.output import write_json

        path = tmp_path / "test.json"
        write_json(sample_listings, path)

        with open(path, encoding="utf-8") as f:
            data = json.load(f)

        assert isinstance(data, list)
        assert len(data) == 2
        assert data[0]["avito_id"] == "100"
        assert data[0]["seller_name"] == "Продавец A"


class TestXLSX:
    def test_xlsx_creates_file(self, tmp_path, sample_listings):
        from src.output.output import write_xlsx

        path = tmp_path / "test.xlsx"
        write_xlsx(sample_listings, None, path)

        assert path.exists()

    def test_xlsx_with_summary(self, tmp_path, sample_listings):
        from src.output.output import write_xlsx
        from src.models import DataOrigin

        summary = RunSummary(
            run_id="test",
            data_origin=DataOrigin.LIVE,
            articles_per_hour=60.0,
            successful_articles_per_hour=54.0,
        )

        path = tmp_path / "test.xlsx"
        write_xlsx(sample_listings, summary, path)

        assert path.exists()
        # Verify it has two sheets
        from openpyxl import load_workbook
        wb = load_workbook(path)
        assert "Listings" in wb.sheetnames
        assert "Run summary" in wb.sheetnames
